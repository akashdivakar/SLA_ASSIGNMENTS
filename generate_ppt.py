import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_output_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    COLOR_BG = RGBColor(15, 23, 42)          # Dark Slate
    COLOR_CARD = RGBColor(30, 41, 59)        # Card Slate
    COLOR_CARD_BORDER = RGBColor(51, 65, 85)
    COLOR_PRIMARY = RGBColor(129, 140, 248)  # Indigo
    COLOR_AMBER = RGBColor(251, 191, 36)     # Amber
    COLOR_TEXT_MAIN = RGBColor(248, 250, 252)# White
    COLOR_TEXT_MUTED = RGBColor(203, 213, 225) # Soft Grey
    COLOR_CODE_BG = RGBColor(10, 15, 29)     # Code Box
    COLOR_CODE_TEXT = RGBColor(56, 189, 248) # Cyan
    COLOR_OUTPUT_GREEN = RGBColor(74, 222, 128) # Green

    slides_content = [
        # Slide 1: Welcome
        {
            "tag": "SEMINAR MASTERCLASS",
            "title": "JavaScript Array Methods vs 'for' Loops (With Outputs)",
            "desc": "Complete side-by-side code comparisons with real execution outputs.",
            "bullets": [
                "✦ Every slide features the exact Input Data, Old Code, New Code, and Output.",
                "✦ See why 1-line modern methods produce identical or better results than 6-line loops.",
                "✦ Covers: map, filter, find, reduce, some, every, and method chaining."
            ],
            "code_title": "Comparison & Output Preview",
            "code": "let nums = [1, 2, 3];\n\n// ❌ OLD: for loop\nlet res1 = [];\nfor (let i = 0; i < nums.length; i++) res1.push(nums[i] * 2);\nconsole.log(res1); // 👉 Output: [2, 4, 6]\n\n// ✅ MODERN: map()\nlet res2 = nums.map(n => n * 2);\nconsole.log(res2); // 👉 Output: [2, 4, 6]",
            "notes": "Welcome everyone! Today every example will show both the traditional for-loop, the modern array method, and the exact output produced by both."
        },
        # Slide 2: MAP
        {
            "tag": "1. MAP METHOD",
            "title": "map() vs 'for' Loop (With Output)",
            "desc": "Goal: Add 5 grace marks to each student in the list.",
            "bullets": [
                "✦ Input: Array of marks [50, 60, 70]",
                "✦ ❌ for Loop: Must initialize empty array and manually push.",
                "✦ ✅ map(): Automatically builds and returns the new array in 1 line.",
                "✦ Output is identical: [55, 65, 75]"
            ],
            "code_title": "map_comparison.js",
            "code": "let marks = [50, 60, 70];\n\n// ❌ OLD WAY (for loop):\nlet res1 = [];\nfor (let i = 0; i < marks.length; i++) {\n    res1.push(marks[i] + 5);\n}\nconsole.log(res1);\n// 👉 Output: [ 55, 65, 75 ]\n\n// ✅ MODERN WAY (map):\nlet res2 = marks.map(m => m + 5);\nconsole.log(res2);\n// 👉 Output: [ 55, 65, 75 ]",
            "notes": "Point out the output: Both produce [55, 65, 75], but map() does it in 1 clean line."
        },
        # Slide 3: FILTER
        {
            "tag": "2. FILTER METHOD",
            "title": "filter() vs 'for' Loop (With Output)",
            "desc": "Goal: Keep only passing marks (35 or above).",
            "bullets": [
                "✦ Input: Array of mixed marks [85, 30, 92, 25, 70]",
                "✦ ❌ for Loop: Needs empty array, loop, and if condition.",
                "✦ ✅ filter(): Automatically collects passing items into an Array.",
                "✦ Output: [85, 92, 70] (Failing marks 30 & 25 removed)"
            ],
            "code_title": "filter_comparison.js",
            "code": "let marks = [85, 30, 92, 25, 70];\n\n// ❌ OLD WAY (for loop):\nlet pass1 = [];\nfor (let i = 0; i < marks.length; i++) {\n    if (marks[i] >= 35) {\n        pass1.push(marks[i]);\n    }\n}\nconsole.log(pass1);\n// 👉 Output: [ 85, 92, 70 ]\n\n// ✅ MODERN WAY (filter):\nlet pass2 = marks.filter(m => m >= 35);\nconsole.log(pass2);\n// 👉 Output: [ 85, 92, 70 ]",
            "notes": "Highlight how filter drops 30 and 25 automatically because the condition returned false."
        },
        # Slide 4: FIND
        {
            "tag": "3. FIND METHOD",
            "title": "find() vs 'for' Loop & filter() (With Output)",
            "desc": "Goal: Find the user with role = 'Admin'.",
            "bullets": [
                "✦ Input: 3 user objects (2 Admins, 1 User)",
                "✦ ❌ for Loop: Needs 'let found = null' and manual 'break'.",
                "✦ ✅ find(): Stops instantly at the 1st match and returns the object.",
                "✦ 💡 vs filter(): find() returns { ... }, filter() returns [ ... ]."
            ],
            "code_title": "find_comparison.js",
            "code": "let users = [\n  { id: 1, role: 'Admin', name: 'Arun' },\n  { id: 2, role: 'User',  name: 'Priya' },\n  { id: 3, role: 'Admin', name: 'Divakar' }\n];\n\n// ❌ OLD WAY (for loop + break):\nlet found1 = null;\nfor (let i = 0; i < users.length; i++) {\n    if (users[i].role === 'Admin') {\n        found1 = users[i];\n        break;\n    }\n}\nconsole.log(found1.name); // 👉 Output: \"Arun\"\n\n// ✅ MODERN WAY (find):\nlet found2 = users.find(u => u.role === 'Admin');\nconsole.log(found2.name); // 👉 Output: \"Arun\"",
            "notes": "Show that find() stops as soon as it sees Arun. It never visits Divakar, saving time."
        },
        # Slide 5: REDUCE
        {
            "tag": "4. REDUCE METHOD",
            "title": "reduce() vs 'for' Loop (With Output)",
            "desc": "Goal: Calculate the total bill of all items in cart.",
            "bullets": [
                "✦ Input: Item prices [100, 250, 50, 200]",
                "✦ ❌ for Loop: Must modify external 'let total = 0' variable.",
                "✦ ✅ reduce(): Encapsulates the running sum inside an accumulator.",
                "✦ Output: ₹600 (Single combined total)"
            ],
            "code_title": "reduce_comparison.js",
            "code": "let prices = [100, 250, 50, 200];\n\n// ❌ OLD WAY (for loop):\nlet total1 = 0;\nfor (let i = 0; i < prices.length; i++) {\n    total1 += prices[i];\n}\nconsole.log('Total: ₹' + total1);\n// 👉 Output: Total: ₹600\n\n// ✅ MODERN WAY (reduce):\nlet total2 = prices.reduce((sum, p) => sum + p, 0);\nconsole.log('Total: ₹' + total2);\n// 👉 Output: Total: ₹600",
            "notes": "Explain how 100 + 250 + 50 + 200 accumulates to 600."
        },
        # Slide 6: SOME
        {
            "tag": "5. SOME METHOD",
            "title": "some() vs 'for' Loop (With Output)",
            "desc": "Goal: Check if ANY student failed (< 35).",
            "bullets": [
                "✦ Input: Marks [80, 90, 40, 95]",
                "✦ ❌ for Loop: Needs flag 'let hasFailed = false' and manual 'break'.",
                "✦ ✅ some(): Returns true / false directly in 1 line.",
                "✦ Output: false (Because all marks are >= 35)"
            ],
            "code_title": "some_comparison.js",
            "code": "let marks = [80, 90, 40, 95];\n\n// ❌ OLD WAY (for loop):\nlet anyFail1 = false;\nfor (let i = 0; i < marks.length; i++) {\n    if (marks[i] < 35) {\n        anyFail1 = true;\n        break;\n    }\n}\nconsole.log(anyFail1); // 👉 Output: false\n\n// ✅ MODERN WAY (some):\nlet anyFail2 = marks.some(m => m < 35);\nconsole.log(anyFail2); // 👉 Output: false",
            "notes": "Ask: What if marks had 25? some() would immediately return true."
        },
        # Slide 7: EVERY
        {
            "tag": "6. EVERY METHOD",
            "title": "every() vs 'for' Loop (With Output)",
            "desc": "Goal: Check if ALL students passed (>= 35).",
            "bullets": [
                "✦ Input: Marks [80, 90, 40, 95]",
                "✦ ❌ for Loop: Needs 'let allPassed = true' and check for failures.",
                "✦ ✅ every(): Directly checks all elements in 1 line.",
                "✦ Output: true (Every student scored >= 35)"
            ],
            "code_title": "every_comparison.js",
            "code": "let marks = [80, 90, 40, 95];\n\n// ❌ OLD WAY (for loop):\nlet allPassed1 = true;\nfor (let i = 0; i < marks.length; i++) {\n    if (marks[i] < 35) {\n        allPassed1 = false;\n        break;\n    }\n}\nconsole.log(allPassed1); // 👉 Output: true\n\n// ✅ MODERN WAY (every):\nlet allPassed2 = marks.every(m => m >= 35);\nconsole.log(allPassed2); // 👉 Output: true",
            "notes": "every() is true only when all elements pass the test."
        },
        # Slide 8: METHOD CHAINING
        {
            "tag": "7. METHOD CHAINING",
            "title": "Method Chaining vs Nested 'for' Loop (With Output)",
            "desc": "Goal: Filter groceries ➔ Add 5% tax ➔ Calculate total bill.",
            "bullets": [
                "✦ Input: 3 items (Rice ₹100, Shirt ₹500, Dal ₹200)",
                "✦ Step 1: Filter groceries (Rice & Dal)",
                "✦ Step 2: Add 5% tax (₹105 & ₹210)",
                "✦ Step 3: Total Bill = ₹105 + ₹210 = ₹315",
                "✦ Output: ₹315 in both, but chaining is 10x cleaner!"
            ],
            "code_title": "chaining_comparison.js",
            "code": "let items = [\n  { name: 'Rice', type: 'grocery', price: 100 },\n  { name: 'Shirt', type: 'cloth',  price: 500 },\n  { name: 'Dal',  type: 'grocery', price: 200 }\n];\n\n// ❌ OLD WAY (for loop):\nlet total1 = 0;\nfor (let i = 0; i < items.length; i++) {\n    if (items[i].type === 'grocery') {\n        total1 += items[i].price * 1.05;\n    }\n}\nconsole.log('Total: ₹' + total1); // 👉 Output: Total: ₹315\n\n// ✅ MODERN WAY (Method Chaining):\nlet total2 = items\n  .filter(i => i.type === 'grocery')\n  .map(i => i.price * 1.05)\n  .reduce((sum, p) => sum + p, 0);\nconsole.log('Total: ₹' + total2); // 👉 Output: Total: ₹315",
            "notes": "Show how chaining turns 3 separate operations into one smooth conveyor belt."
        },
        # Slide 9: SUMMARY TABLE
        {
            "tag": "SUMMARY",
            "title": "Complete Array Methods Master Cheat Sheet",
            "desc": "Quick reference with operation, syntax, and outputs:",
            "bullets": [
                "✦ map() ➔ Changes each item (Input [1,2] ➔ Output [2,4])",
                "✦ filter() ➔ Picks matching items (Input [10,4] ➔ Output [10])",
                "✦ find() ➔ Picks 1st match (Input [{id:1},{id:2}] ➔ Output {id:1})",
                "✦ reduce() ➔ Totals everything (Input [10,20] ➔ Output 30)",
                "✦ some() ➔ Is any matching? (Output: true / false)",
                "✦ every() ➔ Are all matching? (Output: true / false)"
            ],
            "code_title": "Cheat Sheet Matrix",
            "code": "Method    | Purpose          | Return Type    | Sample Output\n--------- | ---------------- | -------------- | -------------\nmap()     | Modify all items | Array (Same len)| [55, 65, 75]\nfilter()  | Pick matching    | Array (Subset)  | [85, 92, 70]\nfind()    | Get 1st match    | Single Item     | { id: 1, ... }\nreduce()  | Total / Combine  | Single Value   | 600\nsome()    | At least one?    | Boolean        | true / false\nevery()   | All matching?    | Boolean        | true / false",
            "notes": "Summarize that modern JavaScript makes our code shorter, cleaner, and less error-prone."
        },
        # Slide 10: Conclusion
        {
            "tag": "CONCLUSION",
            "title": "Thank You & Open Q&A!",
            "desc": "All examples with outputs are ready for your seminar!",
            "bullets": [
                "✦ Check the generated PPT and Interactive HTML file in your project folder.",
                "✦ You are ready to deliver an outstanding presentation!",
                "✦ Questions and live coding practice."
            ],
            "code_title": "Files in Folder",
            "code": "// All Files Ready:\n// 1. JS_Array_Methods_With_Outputs_Seminar.pptx\n// 2. JS_Array_Methods_Seminar_Notes.md\n// 3. JS_Array_Methods_Seminar.html",
            "notes": "Thank everyone for listening and open the floor to questions!"
        }
    ]

    for slide_data in slides_content:
        slide = prs.slides.add_slide(blank_layout)

        # 1. Background Fill
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = COLOR_BG
        bg_shape.line.fill.background()

        # 2. Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        p_tag = tag_box.text_frame.paragraphs[0]
        p_tag.text = slide_data["tag"]
        p_tag.font.size = Pt(12)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_PRIMARY
        p_tag.font.name = "Arial"

        # 3. Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.8))
        p_title = title_box.text_frame.paragraphs[0]
        p_title.text = slide_data["title"]
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_MAIN
        p_title.font.name = "Arial"

        # 4. Description
        desc_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.5))
        p_desc = desc_box.text_frame.paragraphs[0]
        p_desc.text = slide_data["desc"]
        p_desc.font.size = Pt(13)
        p_desc.font.color.rgb = COLOR_TEXT_MUTED
        p_desc.font.name = "Arial"

        # 5. Left Card (Bullets & Input details)
        card_left = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(5.5), Inches(4.7))
        card_left.fill.solid()
        card_left.fill.fore_color.rgb = COLOR_CARD
        card_left.line.color.rgb = COLOR_CARD_BORDER

        bullets_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(5.1), Inches(4.4))
        tf_b = bullets_box.text_frame
        tf_b.word_wrap = True
        for i, bullet in enumerate(slide_data["bullets"]):
            p = tf_b.add_paragraph() if i > 0 else tf_b.paragraphs[0]
            p.text = bullet
            p.font.size = Pt(13)
            p.font.color.rgb = COLOR_TEXT_MAIN
            p.font.name = "Arial"
            p.space_after = Pt(10)

        # 6. Right Card (Side-by-Side Code with OUTPUT)
        card_right = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), Inches(2.2), Inches(5.9), Inches(4.7))
        card_right.fill.solid()
        card_right.fill.fore_color.rgb = COLOR_CODE_BG
        card_right.line.color.rgb = COLOR_PRIMARY

        code_hdr = slide.shapes.add_textbox(Inches(6.8), Inches(2.3), Inches(5.5), Inches(0.4))
        p_ch = code_hdr.text_frame.paragraphs[0]
        p_ch.text = f"💻 {slide_data['code_title']}"
        p_ch.font.size = Pt(12)
        p_ch.font.bold = True
        p_ch.font.color.rgb = COLOR_AMBER
        p_ch.font.name = "Arial"

        code_box = slide.shapes.add_textbox(Inches(6.8), Inches(2.75), Inches(5.5), Inches(4.0))
        tf_c = code_box.text_frame
        tf_c.word_wrap = True
        p_code = tf_c.paragraphs[0]
        p_code.text = slide_data["code"]
        p_code.font.size = Pt(11)
        p_code.font.name = "Courier New"
        p_code.font.color.rgb = COLOR_CODE_TEXT

        # 7. Speaker Notes
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = slide_data["notes"]

    output_path = r"c:\Users\akash\Desktop\SLA_Assignments\JS_Array_Methods_With_Outputs_Seminar.pptx"
    prs.save(output_path)
    print(f"Presentation with Outputs successfully saved to: {output_path}")

if __name__ == "__main__":
    create_output_presentation()
